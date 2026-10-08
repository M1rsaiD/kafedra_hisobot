"""
Генератор отчёта: берёт чистый бланк формы «Umumiy kafedra reytingi» (2026)
и заполняет все 20 листов данными из БД за выбранный ReportingPeriod.

Excel и файлы-подтверждения выгружаются раздельно: Excel — здесь,
ZIP-архив документов по разделам — в archive.py.
"""

import re
from pathlib import Path

import openpyxl
from django.conf import settings
from django.utils.translation import gettext as _

from .layout import BodyWriter, fill_placeholders

TEMPLATE_PATH = Path(__file__).resolve().parent / "templates_store" / "kafedra_hisobot_template.xlsx"


def load_template() -> openpyxl.Workbook:
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(
            f"Файл-шаблон не найден: {TEMPLATE_PATH}. Положите бланк формы "
            "в apps/exporter/templates_store/."
        )
    return openpyxl.load_workbook(TEMPLATE_PATH)


def kafedra_label(department) -> str:
    """«Amaliy matematika va informatika kafedrasi» → «Amaliy matematika va
    informatika» (в шапке формы слово «kafedrasi» уже стоит)."""
    return re.sub(r"\s+kafedrasi\s*$", "", department.name.strip(), flags=re.IGNORECASE)


def safe_filename(text: str) -> str:
    return re.sub(r'[\\/:*?"<>|\s]+', "_", text).strip("_")


def output_basename(period) -> str:
    dept = period.department
    return f"{safe_filename(dept.short_name or kafedra_label(dept))}_{safe_filename(period.period_label)}"


def build_workbook(period):
    """Возвращает (workbook, заполненные листы, предупреждения)."""
    from .sheets import LAYOUTS, SHEET_FILLERS, ReportData

    wb = load_template()
    data = ReportData.load(period)
    context = {
        "kafedra": kafedra_label(period.department),
        "sana": f"{period.period_date:%d.%m.%Y}",
        "yil": period.period_date.year,
        "shtat": period.staff_units if period.staff_units is not None else "______",
    }
    filled = []
    for sheet_name, filler in SHEET_FILLERS.items():
        if sheet_name not in wb.sheetnames:
            data.warnings.append(_("В шаблоне нет листа «%(s)s»") % {"s": sheet_name})
            continue
        ws = wb[sheet_name]
        layout = LAYOUTS[sheet_name]
        fill_placeholders(ws, layout.header_rows, context)
        writer = BodyWriter(ws, layout)
        filler(writer, data)
        writer.finish()
        filled.append(sheet_name)
    warnings = list(dict.fromkeys(data.warnings))  # без повторов, порядок сохраняем
    return wb, filled, warnings


def export_report(department, period):
    """
    Собирает итоговый .xlsx для кафедры за отчётный период и сохраняет в
    EXPORTS_DIR. Возвращает (путь, заполненные листы, предупреждения).
    """
    wb, filled, warnings = build_workbook(period)
    path = Path(settings.EXPORTS_DIR) / f"{output_basename(period)}.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path, filled, warnings
