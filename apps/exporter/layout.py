"""
Перестройка «тела» листа формы под реальное число строк.

Почему не просто «писать в ячейки шаблона»: в официальной форме под каждый
лист заготовлено фиксированное число строк, а людей/статей на кафедре может
быть и больше, и меньше; к тому же внизу каждого листа идут примечание
«*… ilova qilinadi» и подпись «Rektor» в объединённых ячейках. openpyxl не
сдвигает объединённые ячейки при insert_rows/delete_rows, поэтому тело листа
собирается заново: шапка остаётся как в форме, строки-прототипы (стиль и
объединения) берутся из шаблона, подвал переносится под последнюю строку.

Контракт чистого шаблона (templates_store/kafedra_hisobot_template.xlsx) —
для каждого листа сверху вниз:

    1..H            шапка формы (как у министерства), H = SheetLayout.header_rows
    H+1             строка-прототип заголовка секции (только если has_marker)
    следующая       строка-прототип данных
    следующая       строка-прототип итога «Jami» (только если has_total)
    остальное       подвал: примечания, «Rektor ... (imzo) ... (f.i.sh.)», «Sana»

В шапке допускаются подстановки {kafedra}, {sana}, {yil}, {shtat}.
"""

from copy import copy
from dataclasses import dataclass, field

from openpyxl.utils import get_column_letter


@dataclass
class SheetLayout:
    header_rows: int
    ncols: int
    has_marker: bool = False
    marker_col: int = 1
    has_total: bool = True

    def proto_rows(self):
        row = self.header_rows + 1
        marker = data = total = None
        if self.has_marker:
            marker, row = row, row + 1
        data, row = row, row + 1
        if self.has_total:
            total, row = row, row + 1
        return marker, data, total, row  # row = начало подвала


@dataclass
class RowSnapshot:
    """Значения, стили, высота и объединения одной строки (объединения
    хранятся как пары колонок, т.к. строка будет вставлена в другом месте)."""

    values: dict = field(default_factory=dict)
    styles: dict = field(default_factory=dict)
    height: float | None = None
    merges: list = field(default_factory=list)   # [(min_col, max_col, row_span)]


def _merges_starting_at(ws, row):
    out = []
    for rng in ws.merged_cells.ranges:
        if rng.min_row == row:
            out.append((rng.min_col, rng.max_col, rng.max_row - rng.min_row + 1))
    return out


def snapshot_row(ws, row, ncols):
    snap = RowSnapshot(height=ws.row_dimensions[row].height if row in ws.row_dimensions else None)
    max_col = max([ncols] + [c.column for c in ws[row] if c.value is not None])
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        if cell.has_style:
            snap.styles[col] = copy(cell._style)
        if cell.value is not None:
            snap.values[col] = cell.value
    snap.merges = _merges_starting_at(ws, row)
    return snap


def paste_row(ws, row, snap, values=None, keep_values=False):
    for col, style in snap.styles.items():
        ws.cell(row=row, column=col)._style = copy(style)
    if keep_values:
        for col, value in snap.values.items():
            ws.cell(row=row, column=col).value = value
    for col, value in (values or {}).items():
        ws.cell(row=row, column=col).value = value
    if snap.height:
        ws.row_dimensions[row].height = snap.height
    for min_col, max_col, span in snap.merges:
        if max_col > min_col or span > 1:
            ws.merge_cells(
                f"{get_column_letter(min_col)}{row}:{get_column_letter(max_col)}{row + span - 1}"
            )


def clear_below(ws, last_kept_row):
    """Убирает всё ниже шапки: объединения, значения, стили, высоты строк."""
    for rng in list(ws.merged_cells.ranges):
        if rng.max_row > last_kept_row:
            ws.unmerge_cells(str(rng))
    if ws.max_row > last_kept_row:
        ws.delete_rows(last_kept_row + 1, ws.max_row - last_kept_row)
    for r in [r for r in ws.row_dimensions if r > last_kept_row]:
        del ws.row_dimensions[r]


class BodyWriter:
    """Собирает тело листа: секции, строки данных, итог, затем подвал."""

    def __init__(self, ws, layout: SheetLayout):
        self.ws = ws
        self.layout = layout
        marker, data, total, footer_start = layout.proto_rows()
        n = layout.ncols
        self.marker_proto = snapshot_row(ws, marker, n) if marker else None
        self.data_proto = snapshot_row(ws, data, n)
        self.total_proto = snapshot_row(ws, total, n) if total else None
        self.footer = [snapshot_row(ws, r, n) for r in range(footer_start, ws.max_row + 1)]
        clear_below(ws, layout.header_rows)
        self.row = layout.header_rows + 1

    def marker(self, text):
        proto = self.marker_proto or self.data_proto
        paste_row(self.ws, self.row, proto, {self.layout.marker_col: text})
        self.ws.row_dimensions[self.row].height = None  # длинные заголовки секций листа 5
        self.row += 1

    def data(self, values):
        paste_row(self.ws, self.row, self.data_proto, values)
        # высоту строк данных не фиксируем — Excel подберёт по переносу текста
        self.ws.row_dimensions[self.row].height = None
        self.row += 1

    def total(self, values):
        paste_row(self.ws, self.row, self.total_proto or self.data_proto, values)
        self.row += 1

    def finish(self):
        for snap in self.footer:
            paste_row(self.ws, self.row, snap, keep_values=True)
            self.row += 1


def fill_placeholders(ws, header_rows, context):
    """{kafedra}, {sana} и т.п. в шапке листа."""
    for row in ws.iter_rows(min_row=1, max_row=header_rows):
        for cell in row:
            if isinstance(cell.value, str) and "{" in cell.value:
                try:
                    cell.value = cell.value.format(**context)
                except (KeyError, IndexError, ValueError):
                    pass
