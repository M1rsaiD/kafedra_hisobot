import datetime
import io
import shutil
import tempfile
import zipfile
from decimal import Decimal

import openpyxl
from django.core.files.base import ContentFile
from django.test import TestCase, override_settings

from apps.core.models import Department, Faculty, University
from apps.dictionaries.models import AcademicDegree, EmploymentType, Position
from apps.exporter.archive import export_archive
from apps.exporter.engine import build_workbook, export_report
from apps.exporter.importer import Importer, normalize_name
from apps.reports.models import (
    Funding, Patent, PatentAuthor, Publication, PublicationAuthor, TeacherHIndex,
)
from apps.staff.models import Teacher, TeacherPeriodSnapshot
from apps.workflow.models import ReportingPeriod

TMP = tempfile.mkdtemp()


def make_department(name="Amaliy matematika va informatika kafedrasi"):
    uni = University.objects.create(name="Farg‘ona davlat universiteti")
    fac = Faculty.objects.create(university=uni, name="Matematika-informatika")
    return Department.objects.create(faculty=fac, name=name)


def employment(code):
    order = {"asosiy": 1, "ichki_orindosh": 2, "tashqi_orindosh": 3, "soatbay": 4}[code]
    return EmploymentType.objects.get_or_create(code=code, defaults={"name": code, "sort_order": order})[0]


def add_teacher(period, name, code="asosiy", degree=None, **fields):
    t = Teacher.objects.create(department=period.department, full_name=name, **fields)
    TeacherPeriodSnapshot.objects.create(
        teacher=t, reporting_period=period, employment_type=employment(code),
        position=Position.objects.get_or_create(name="dotsent")[0], stavka=Decimal("1.5"),
        academic_degree=AcademicDegree.objects.get_or_create(code=degree, defaults={"name": degree.upper()})[0]
        if degree else None,
    )
    return t


def sheet_rows(ws, first_row):
    """{ФИО: строка значений} для строк с номером в A."""
    out = {}
    for row in ws.iter_rows(min_row=first_row, values_only=True):
        if isinstance(row[0], int) and row[1]:
            out[row[1]] = row
    return out


@override_settings(MEDIA_ROOT=TMP + "/media", EXPORTS_DIR=TMP + "/exports")
class ExportTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(TMP, ignore_errors=True)

    def setUp(self):
        self.dept = make_department()
        self.period = ReportingPeriod.objects.create(
            department=self.dept, period_label="2026-iyun", period_date=datetime.date(2026, 6, 25))
        self.a = add_teacher(self.period, "Aliyev Anvar", degree="phd", birth_date=datetime.date(1980, 5, 1))
        self.b = add_teacher(self.period, "Botirov Bobur", degree="dsc", birth_date=datetime.date(1970, 1, 1))
        self.c = add_teacher(self.period, "Valiyev Vali", code="ichki_orindosh")

    def publication(self, authors, total, **kw):
        pub = Publication.objects.create(teacher=authors[0], reporting_period=self.period,
                                         authors_total=total, publish_year=2026, **kw)
        for t in authors:
            PublicationAuthor.objects.create(publication=pub, teacher=t)
        return pub

    def test_article_shares_split_by_total_authors(self):
        # 2 автора кафедры, всего 2 автора → по 0.5 в Q1
        self.publication([self.a, self.b], 2, quartile="Q1", article_title="Shared")
        # 1 автор кафедры из 3 авторов → 0.33 в Q2
        self.publication([self.a], 3, quartile="Q2", article_title="Three authors")
        # статья прошлого года в 2,1 не идёт, но в 2,2 есть
        old = self.publication([self.b], 1, quartile="Q1", article_title="Old")
        Publication.objects.filter(pk=old.pk).update(publish_year=2020)

        wb, _, warnings = build_workbook(self.period)
        rows = sheet_rows(wb["2,1"], 3)
        self.assertEqual(rows["Aliyev Anvar"][4], 0.5)    # E = Q1
        self.assertEqual(rows["Aliyev Anvar"][5], 0.33)   # F = Q2
        self.assertEqual(rows["Botirov Bobur"][4], 0.5)
        self.assertEqual(warnings, [])
        titles = [r[4] for r in wb["2,2"].iter_rows(min_row=3, values_only=True) if isinstance(r[0], int)]
        self.assertEqual(sorted(titles), ["Old", "Shared", "Shared", "Three authors"])

    def test_patent_share(self):
        p = Patent.objects.create(reporting_period=self.period, title="P", patent_type="invention",
                                  authors_total=2, has_license_agreement=True)
        PatentAuthor.objects.create(patent=p, teacher=self.a)
        dgu = Patent.objects.create(reporting_period=self.period, title="DGU",
                                    patent_type=Patent.PatentType.SOFTWARE, authors_total=1)
        PatentAuthor.objects.create(patent=dgu, teacher=self.a)
        wb, _, _ = build_workbook(self.period)
        rows = sheet_rows(wb["2,4"], 3)
        self.assertEqual(rows["Aliyev Anvar"][2:4], (0.5, 0.5))   # свидетельство DGU в 2,4 не считается

    def test_any_number_of_rows_and_footer_moves_down(self):
        for i in range(40):
            add_teacher(self.period, f"Xodim {i:02d}")
        wb, filled, _ = build_workbook(self.period)
        self.assertEqual(len(filled), 20)
        ws = wb["5"]
        names = [r[1] for r in ws.iter_rows(min_row=7, values_only=True) if isinstance(r[0], int)]
        self.assertEqual(len(names), 43)
        self.assertEqual(names[:2], ["Aliyev Anvar", "Botirov Bobur"])
        # подпись ректора — ниже последней строки, объединённые ячейки не пересекаются
        rektor_row = next(c.row for c in ws["A"] if isinstance(c.value, str) and "Rektor" in c.value)
        last_data_row = max(c.row for c in ws["B"] if c.value == "Xodim 39")
        self.assertGreater(rektor_row, last_data_row)
        self.assertIn("Amaliy matematika va informatika kafedrasi", ws["A2"].value)
        self.assertIn("25.06.2026", ws["A2"].value)

    def test_sections_and_totals(self):
        TeacherHIndex.objects.create(teacher=self.a, reporting_period=self.period, h_index=3)
        TeacherHIndex.objects.create(teacher=self.c, reporting_period=self.period, h_index=2)
        Funding.objects.create(teacher=self.b, reporting_period=self.period,
                               source_type="business_contract", amount=Decimal("1500000"))
        wb, _, _ = build_workbook(self.period)
        ws = wb["2,3"]
        values = [r for r in ws.iter_rows(min_row=3, max_col=3, values_only=True)]
        self.assertEqual(values[0][0], "Asosiy shtatdagilar")
        self.assertIn(("Ichki o‘rindosh", None, None), values)
        self.assertIn((None, "Jami:", 5), values)
        self.assertEqual(sheet_rows(wb["2,6"], 3)["Botirov Bobur"][4], 1500000)
        ages = sheet_rows(wb["1,2"], 3)
        self.assertEqual(ages["Aliyev Anvar"][3], 46)

    def test_records_of_teachers_outside_staffing_are_reported(self):
        outsider = Teacher.objects.create(department=self.dept, full_name="Tashqi Toir")
        self.publication([outsider], 1, quartile="Q1", article_title="X")
        _, _, warnings = build_workbook(self.period)
        self.assertTrue(any("Tashqi Toir" in w for w in warnings))

    def test_archive_contains_files_by_section_and_index(self):
        pub = self.publication([self.a, self.b], 2, article_title="Shared/Article: one")
        pub.article_file.save("a.pdf", ContentFile(b"%PDF-1.4 test"))
        self.a.diploma_file.save("d.pdf", ContentFile(b"%PDF diploma"))
        path, stats = export_archive(self.period)
        with zipfile.ZipFile(path) as zf:
            names = zf.namelist()
            index = openpyxl.load_workbook(io.BytesIO(zf.read("_fayllar_royxati.xlsx")))
        self.assertIn("2,1-2,2 Maqolalar/Aliyev, Botirov - 2026 - Shared Article one.pdf", names)
        self.assertIn("1,1 Ilmiy daraja diplomlari/Aliyev Anvar - diplom.pdf", names)
        self.assertEqual(stats["files"], 2)
        # нет диплома у Ботирова и приказов о приёме у всех троих
        self.assertEqual(stats["missing"], 4)
        statuses = [r[4] for r in index.active.iter_rows(min_row=2, values_only=True)]
        self.assertEqual(statuses.count("YO‘Q / НЕТ"), 4)

    def test_import_roundtrip(self):
        """Экспорт → импорт в новый период даёт те же штат, h-индексы, статьи."""
        self.a.diploma_series, self.a.diploma_number = "01", "000123"
        self.a.save()
        TeacherHIndex.objects.create(teacher=self.a, reporting_period=self.period, h_index=4)
        self.publication([self.a, self.b], 2, article_title="Shared", citations_count=7)
        path, _, _ = export_report(self.dept, self.period)

        new_period = ReportingPeriod.objects.create(
            department=self.dept, period_label="2026-dekabr", period_date=datetime.date(2026, 12, 31))
        stats, unmatched = Importer(path, new_period).run()
        self.assertEqual(unmatched, [])
        self.assertEqual(new_period.teacher_snapshots.count(), 3)
        self.assertEqual(TeacherHIndex.objects.get(reporting_period=new_period).h_index, 4)
        pub = Publication.objects.get(reporting_period=new_period)
        self.assertEqual(pub.citations_count, 7)
        self.assertEqual({t.full_name for t in pub.authors.all()}, {"Aliyev Anvar", "Botirov Bobur"})

    def test_name_normalization_matches_cyrillic_and_variants(self):
        self.assertEqual(normalize_name("Каримова Дилноза Хабибуллаевна"),
                         normalize_name("Karimova Dilnoza Xabibullayevna"))
        self.assertEqual(normalize_name("Турсункулов Максаджон"), normalize_name("Tursunqulov Maqsadjon"))
