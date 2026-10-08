"""
ZIP-архив подтверждающих документов отчёта — отдельно от Excel.

Структура архива (папка = раздел формы):

    1,1 Ilmiy daraja diplomlari/<F.I.Sh.> - diplom.pdf
    1,1 Ishga qabul buyruqlari/<F.I.Sh.> - buyruq.pdf
    1,3 Xorijiy magistr diplomlari/...
    2,1-2,2 Maqolalar/<Mualliflar> - <yil> - <nomi>.pdf
    2,3 h-indeks/...
    2,4-2,5 Patentlar va guvohnomalar/...
    2,6 Jalb etilgan mablag'lar/...
    2,7 Himoyalar/...
    3,1 Xalqaro reputatsiya/...
    3,3 Talabalar mobilligi/...
    3,4 PO' xalqaro mobilligi/...
    3,5 Qo'shma ta'lim dasturlari/...
    4 Bitiruvchilar/...
    _fayllar_royxati.xlsx   — реестр: что приложено и чего не хватает

collect_documents() используется и архивом, и панелью завкафедрой
(сколько документов не загружено у каждого преподавателя).
"""

import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from django.conf import settings
from openpyxl.styles import Font, PatternFill

from apps.reports.models import (
    Publication, TeacherHIndex, Patent, Funding, PostgradDefense, ReputationVote,
    TeacherMobility, StudentMobility, GraduateEmployment,
)
from apps.staff.models import TeacherPeriodSnapshot

from .engine import output_basename

SEC_DIPLOMA = "1,1 Ilmiy daraja diplomlari"
SEC_HIRE = "1,1 Ishga qabul buyruqlari"
SEC_MASTER = "1,3 Xorijiy magistr diplomlari"
SEC_PUBS = "2,1-2,2 Maqolalar"
SEC_HINDEX = "2,3 h-indeks"
SEC_PATENTS = "2,4-2,5 Patentlar va guvohnomalar"
SEC_FUNDING = "2,6 Jalb etilgan mablag'lar"
SEC_DEFENSE = "2,7 Himoyalar"
SEC_REPUTATION = "3,1 Xalqaro reputatsiya"
SEC_STUDENT_MOB = "3,3 Talabalar mobilligi"
SEC_TEACHER_MOB = "3,4 PO' xalqaro mobilligi"
SEC_JOINT = "3,5 Qo'shma ta'lim dasturlari"
SEC_GRADUATES = "4 Bitiruvchilar"


@dataclass
class DocumentEntry:
    section: str
    owner: str                # чей документ (ФИО или авторы)
    title: str                # что за запись
    file: object              # FieldFile или None, если не загружен
    teacher_ids: tuple = ()   # кому из преподавателей «засчитать» отсутствие файла

    @property
    def present(self):
        return bool(self.file)

    def archive_name(self):
        ext = Path(self.file.name).suffix.lower() if self.file else ""
        stem = _clean(f"{self.owner} - {self.title}")[:150].rstrip(" .-")
        return f"{self.section}/{stem}{ext}"


def _clean(text):
    text = re.sub(r'[\\/:*?"<>|\r\n\t]+', " ", str(text))
    return re.sub(r"\s+", " ", text).strip()


def _authors_label(obj):
    names = [t.full_name.split()[0] for t in obj.authors.all()]
    return ", ".join(names) or "—"


def collect_documents(period):
    docs = []
    snaps = (TeacherPeriodSnapshot.objects.filter(reporting_period=period,
                                                  employment_type__code__in=("asosiy", "ichki_orindosh"))
             .select_related("teacher").order_by("teacher__full_name"))
    for snap in snaps:
        t = snap.teacher
        if snap.academic_degree_id or snap.academic_title_id:
            docs.append(DocumentEntry(SEC_DIPLOMA, t.full_name, "diplom", t.diploma_file, (t.id,)))
        docs.append(DocumentEntry(SEC_HIRE, t.full_name, "buyruq", t.hire_order_file, (t.id,)))
        if t.foreign_master_rank or t.foreign_master_university:
            docs.append(DocumentEntry(SEC_MASTER, t.full_name, "magistr diplomi", t.master_diploma_file, (t.id,)))

    for pub in Publication.objects.filter(reporting_period=period).prefetch_related("authors"):
        title = f"{pub.publish_year or ''} - {pub.article_title or pub.journal_name or pub.pk}"
        docs.append(DocumentEntry(SEC_PUBS, _authors_label(pub), title, pub.article_file,
                                  tuple(a.id for a in pub.authors.all())))

    for rec in TeacherHIndex.objects.filter(reporting_period=period).select_related("teacher"):
        docs.append(DocumentEntry(SEC_HINDEX, rec.teacher.full_name, f"h-index {rec.h_index}",
                                  rec.document_file, (rec.teacher_id,)))

    for p in Patent.objects.filter(reporting_period=period).prefetch_related("authors"):
        docs.append(DocumentEntry(SEC_PATENTS, _authors_label(p), f"{p.number or p.pk} - {p.title}",
                                  p.document_file, tuple(a.id for a in p.authors.all())))

    for f in Funding.objects.filter(reporting_period=period).select_related("teacher"):
        docs.append(DocumentEntry(SEC_FUNDING, f.teacher.full_name,
                                  f"{f.get_source_type_display()} - {f.amount:.0f}", f.document_file,
                                  (f.teacher_id,)))

    for d in PostgradDefense.objects.filter(reporting_period=period).select_related("advisor_teacher"):
        docs.append(DocumentEntry(SEC_DEFENSE, d.advisor_teacher.full_name,
                                  f"{d.degree_type} - {d.defended_full_name}", d.document_file,
                                  (d.advisor_teacher_id,)))

    for v in ReputationVote.objects.filter(reporting_period=period).select_related("teacher"):
        docs.append(DocumentEntry(SEC_REPUTATION, v.teacher.full_name,
                                  f"{v.get_source_display()} - {v.votes_count}", v.document_file,
                                  (v.teacher_id,)))

    for m in TeacherMobility.objects.filter(reporting_period=period).select_related("teacher"):
        docs.append(DocumentEntry(SEC_TEACHER_MOB, m.teacher.full_name,
                                  f"{m.direction} - {m.partner_institution_name or m.pk}", m.document_file,
                                  (m.teacher_id,)))

    for m in StudentMobility.objects.filter(reporting_period=period).select_related("student"):
        section = SEC_JOINT if m.program_type == StudentMobility.ProgramType.JOINT_PROGRAM else SEC_STUDENT_MOB
        docs.append(DocumentEntry(section, m.student.full_name,
                                  f"{m.get_partner_rank_display()} - {m.partner_institution_name or m.pk}",
                                  m.document_file))

    for e in GraduateEmployment.objects.filter(reporting_period=period).select_related("graduate"):
        docs.append(DocumentEntry(SEC_GRADUATES, e.graduate.full_name, e.get_status_display(),
                                  e.document_file))
    return docs


def _write_index(docs, names):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fayllar"
    ws.append(["Bo‘lim / Раздел", "Kimniki / Чей", "Yozuv / Запись", "Fayl / Файл", "Holat / Статус"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    missing_fill = PatternFill("solid", fgColor="FBE6E6")
    for doc, name in zip(docs, names):
        status = "bor / есть" if name else "YO‘Q / НЕТ"
        ws.append([doc.section, doc.owner, doc.title, name or "", status])
        row = ws.max_row
        if name:
            ws.cell(row=row, column=4).hyperlink = name
        else:
            for col in range(1, 6):
                ws.cell(row=row, column=col).fill = missing_fill
    for col, width in zip("ABCDE", (34, 36, 50, 70, 14)):
        ws.column_dimensions[col].width = width
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    return wb


def export_archive(period):
    """Собирает ZIP со всеми загруженными файлами периода. Возвращает
    (путь, {"files": N, "missing": M})."""
    docs = collect_documents(period)
    path = Path(settings.EXPORTS_DIR) / f"{output_basename(period)}_fayllar.zip"
    path.parent.mkdir(parents=True, exist_ok=True)

    used, names, files = set(), [], 0
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for doc in docs:
            name = None
            if doc.present:
                name = doc.archive_name()
                base, ext = name.rsplit(".", 1) if "." in Path(name).name else (name, "")
                n = 2
                while name in used:
                    name = f"{base} ({n}).{ext}" if ext else f"{base} ({n})"
                    n += 1
                try:
                    with doc.file.open("rb") as fh:
                        zf.writestr(name, fh.read())
                except (FileNotFoundError, OSError):
                    name = None  # запись есть, а файла на диске нет — считаем отсутствующим
                else:
                    used.add(name)
                    files += 1
            names.append(name)

        index = _write_index(docs, names)
        tmp = path.with_suffix(".index.xlsx")
        index.save(tmp)
        zf.write(tmp, "_fayllar_royxati.xlsx")
        tmp.unlink()

    return path, {"files": files, "missing": sum(1 for n in names if not n)}
