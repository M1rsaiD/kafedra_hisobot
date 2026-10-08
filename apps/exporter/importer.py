"""
Импорт уже заполненного вручную файла формы 2026 в БД — чтобы перейти на
платформу без повторного набора: преподаватели потом только проверяют свои
записи и прикладывают файлы.

Что переносится:
    5    — преподаватели, тип занятости, должность, ставка, степень, звание
    1,1  — диплом (серия/номер), специальность, приказ о приёме, зарубежный ОТМ
    1,2  — точная дата рождения
    2,2  — статьи (одна статья у нескольких преподавателей → одна запись с соавторами)
    2,3  — h-индекс
    2,4  — патенты (в форме только количество → записи-заготовки «уточните»)
    2,6  — привлечённые средства
    3,1  — голоса за репутацию
    3,3  — мобильность студентов
    3,4  — мобильность ППС (в форме только количество → записи-заготовки)

Лист 2,1 (квартили) не переносится: в форме это свод по человеку, а не по
статье, — квартиль каждой статьи преподаватель указывает сам, свод
пересчитается автоматически с учётом долей 1/N.

ФИО в разных листах одного файла пишутся по-разному (кириллица/латиница,
«Ortiqovich»/«Ortikovich»), поэтому сопоставление нечёткое — по
транслитерированной и нормализованной форме фамилии и имени.
"""

import datetime
import difflib
import re
from collections import defaultdict
from decimal import Decimal

import openpyxl
from django.db import transaction

from apps.dictionaries.models import (
    AcademicDegree, AcademicTitle, EmploymentType, ForeignInstitutionRank as Rank, Position,
)
from apps.reports.models import (
    Funding, Patent, PatentAuthor, Publication, PublicationAuthor, ReputationVote,
    StudentMobility, TeacherHIndex, TeacherMobility,
)
from apps.staff.models import Teacher, TeacherPeriodSnapshot
from apps.students.models import Student

IMPORT_NOTE = "Exceldan import — aniqlang / импорт из Excel — уточните"

CYR = {
    "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "yo", "ж": "j", "з": "z",
    "и": "i", "й": "y", "к": "k", "л": "l", "м": "m", "н": "n", "о": "o", "п": "p", "р": "r",
    "с": "s", "т": "t", "у": "u", "ф": "f", "х": "x", "ц": "s", "ч": "ch", "ш": "sh", "щ": "sh",
    "ъ": "", "ы": "i", "ь": "", "э": "e", "ю": "yu", "я": "ya", "ў": "o", "қ": "q", "ғ": "g",
    "ҳ": "h", "ә": "a", "ө": "o", "ү": "u", "ұ": "u", "ң": "n", "і": "i",
}


def normalize_name(name):
    """Ключ для нечёткого сравнения: латиница, без апострофов, k/q, h/x и т.п."""
    s = " ".join(str(name or "").split()).lower()
    s = "".join(CYR.get(ch, ch) for ch in s)
    s = re.sub(r"[‘’ʻʼ'`´]", "", s)
    for a, b in (("q", "k"), ("x", "h"), ("yo", "o"), ("ye", "e"), ("iy", "i"), ("y", "i")):
        s = s.replace(a, b)
    s = re.sub(r"[^a-z ]", "", s)
    return " ".join(s.split()[:2])  # фамилия + имя достаточно и устойчивее


def clean_text(value):
    if value is None:
        return ""
    return " ".join(str(value).split())


def to_number(value):
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value).replace(",", ".").replace(" ", ""))
    except Exception:
        return None


class Roster:
    def __init__(self, department):
        self.department = department
        self.by_key = {normalize_name(t.full_name): t for t in Teacher.objects.filter(department=department)}

    def add(self, teacher):
        self.by_key[normalize_name(teacher.full_name)] = teacher

    def find(self, name):
        key = normalize_name(name)
        if not key:
            return None
        if key in self.by_key:
            return self.by_key[key]
        match = difflib.get_close_matches(key, list(self.by_key), n=1, cutoff=0.82)
        return self.by_key[match[0]] if match else None


def iter_rows(ws, start_row, ncols):
    """Строки данных формы: номер в A и имя в B (заголовки секций и итог
    «Jami»/«O‘rtacha» номера не имеют); останавливаемся на подвале
    (Rektor / примечание со звёздочкой)."""
    for row in ws.iter_rows(min_row=start_row, max_col=ncols, values_only=True):
        a, b = row[0], row[1] if len(row) > 1 else None
        text = clean_text(a) or clean_text(b)
        if text.startswith(("Rektor", "*")) or clean_text(b).startswith("Jami"):
            break
        if b and clean_text(b) and isinstance(a, (int, float)):
            yield row


class Importer:
    def __init__(self, path, period, log=print):
        self.wb = openpyxl.load_workbook(path, data_only=True)
        self.period = period
        self.department = period.department
        self.log = log
        self.roster = Roster(self.department)
        self.stats = defaultdict(int)
        self.unmatched = []

    def teacher(self, name, sheet):
        t = self.roster.find(name)
        if t is None:
            self.unmatched.append(f"{sheet}: «{clean_text(name)}» не найден в листе 5")
        return t

    @transaction.atomic
    def run(self):
        self.sheet_5()
        self.sheet_1_1()
        self.sheet_1_2()
        self.sheet_2_2()
        self.sheet_2_3()
        self.sheet_2_4()
        self.sheet_2_6()
        self.sheet_3_1()
        self.sheet_3_3()
        self.sheet_3_4()
        return self.stats, self.unmatched

    # ----------------------------------------------------------------- 5
    def sheet_5(self):
        ws = self.wb["5"]
        types = {
            "asosiy": ("asosiy", "Asosiy shtat", 1),
            "ichki": ("ichki_orindosh", "Ichki o‘rindosh", 2),
            "tashqi": ("tashqi_orindosh", "Tashqi o‘rindosh", 3),
            "soatbay": ("soatbay", "Soatbay", 4),
        }
        current = None
        for row in ws.iter_rows(min_row=6, max_col=7, values_only=True):
            a, b, c, d, e, f, g = row
            label = clean_text(b).lower()
            if b and not isinstance(a, (int, float)):
                for prefix, (code, name, order) in types.items():
                    if label.startswith(prefix):
                        current, _ = EmploymentType.objects.get_or_create(
                            code=code, defaults={"name": name, "sort_order": order})
                continue
            if clean_text(a).startswith(("Rektor", "…")) or current is None or not clean_text(b):
                continue
            name = clean_text(b)
            teacher = self.roster.find(name)
            if teacher is None:
                teacher = Teacher.objects.create(department=self.department, full_name=name)
                self.roster.add(teacher)
                self.stats["преподавателей создано"] += 1
            if c and not teacher.birth_date and isinstance(c, (int, float)):
                # точная дата придёт из листа 1,2 (если человек там есть)
                teacher.birth_date = datetime.date(int(c), 1, 1)
                teacher.save(update_fields=["birth_date"])
            TeacherPeriodSnapshot.objects.update_or_create(
                teacher=teacher, reporting_period=self.period,
                defaults={
                    "employment_type": current,
                    "position": Position.objects.get_or_create(name=clean_text(d))[0] if d else None,
                    "stavka": to_number(e) or Decimal(0),
                    "academic_degree": self.degree(f),
                    "academic_title": self.title(g),
                },
            )
            self.stats["записей штатной ведомости"] += 1

    @staticmethod
    def degree(value):
        key = clean_text(value).lower()
        if not key:
            return None
        code = {"phd": "phd", "dsc": "dsc", "fan nomzodi": "fan_nomzodi",
                "fan doktori": "fan_doktori", "akademik": "akademik"}.get(key)
        if not code:
            return None
        names = {"phd": "PhD", "dsc": "DSc", "fan_nomzodi": "Fan nomzodi",
                 "fan_doktori": "Fan doktori", "akademik": "Akademik"}
        return AcademicDegree.objects.get_or_create(code=code, defaults={"name": names[code]})[0]

    @staticmethod
    def title(value):
        key = clean_text(value).lower()
        if not key:
            return None
        if "prof" in key:
            return AcademicTitle.objects.get_or_create(name="Professor")[0]
        if "dots" in key:
            return AcademicTitle.objects.get_or_create(name="Dotsent")[0]
        return AcademicTitle.objects.get_or_create(name=clean_text(value))[0]

    # ----------------------------------------------------------------- 1,1
    def sheet_1_1(self):
        for row in iter_rows(self.wb["1,1"], 4, 11):
            t = self.teacher(row[1], "1,1")
            if not t:
                continue
            t.foreign_university_name = clean_text(row[2]) or t.foreign_university_name
            for s_col, n_col in ((3, 4), (5, 6), (7, 8)):
                if row[s_col] or row[n_col]:
                    t.diploma_series = clean_text(row[s_col])
                    t.diploma_number = clean_text(row[n_col])
            t.specialty_name = clean_text(row[9]) or t.specialty_name
            order = clean_text(row[10])
            if order:
                number, _, rest = order.partition(",")
                t.hire_order_number = number.strip()
                m = re.search(r"(\d{1,2})\.(\d{1,2})\.(\d{4})", rest or order)
                if m:
                    d, mo, y = map(int, m.groups())
                    t.hire_order_date = datetime.date(y, mo, d)
            t.save()
            self.stats["дипломов/приказов (1,1)"] += 1

    # ----------------------------------------------------------------- 1,2
    def sheet_1_2(self):
        for row in iter_rows(self.wb["1,2"], 3, 4):
            t = self.teacher(row[1], "1,2")
            if t and isinstance(row[2], (datetime.date, datetime.datetime)):
                t.birth_date = row[2].date() if isinstance(row[2], datetime.datetime) else row[2]
                t.save(update_fields=["birth_date"])
                self.stats["дат рождения (1,2)"] += 1

    # ----------------------------------------------------------------- 2,2
    def sheet_2_2(self):
        by_title = {}
        for row in iter_rows(self.wb["2,2"], 3, 8):
            t = self.teacher(row[1], "2,2")
            if not t:
                continue
            title = clean_text(row[4])
            key = re.sub(r"[^a-zа-я0-9]", "", title.lower())
            pub = by_title.get(key)
            if pub is None:
                year, month = None, None
                m = re.match(r"(\d{4})(?:\D+(\d{1,2}))?", clean_text(row[3]).replace(".0", ""))
                if m:
                    year = int(m.group(1))
                    month = int(m.group(2)) if m.group(2) and 1 <= int(m.group(2)) <= 12 else None
                url = clean_text(row[6])
                pub = Publication.objects.create(
                    teacher=t, reporting_period=self.period,
                    journal_name=clean_text(row[2])[:255], publish_year=year, publish_month=month,
                    article_title=title[:500], language=clean_text(row[5])[:30],
                    scopus_url=url[:500] if url.startswith("http") else "",
                    citations_count=int(to_number(row[7]) or 0),
                )
                by_title[key] = pub
                self.stats["статей (2,2)"] += 1
            PublicationAuthor.objects.get_or_create(publication=pub, teacher=t)

    # ----------------------------------------------------------------- 2,3
    def sheet_2_3(self):
        for row in iter_rows(self.wb["2,3"], 3, 3):
            t = self.teacher(row[1], "2,3")
            h = to_number(row[2])
            if t and h is not None:
                TeacherHIndex.objects.update_or_create(
                    teacher=t, reporting_period=self.period, defaults={"h_index": int(h)})
                self.stats["h-индексов (2,3)"] += 1

    # ----------------------------------------------------------------- 2,4
    def sheet_2_4(self):
        for row in iter_rows(self.wb["2,4"], 3, 4):
            t = self.teacher(row[1], "2,4")
            count = int(round(to_number(row[2]) or 0))
            licensed = int(round(to_number(row[3]) or 0))
            for i in range(count if t else 0):
                patent = Patent.objects.create(
                    reporting_period=self.period, title=IMPORT_NOTE,
                    patent_type=Patent.PatentType.INVENTION, has_license_agreement=i < licensed,
                )
                PatentAuthor.objects.create(patent=patent, teacher=t)
                self.stats["патентов-заготовок (2,4)"] += 1

    # ----------------------------------------------------------------- 2,6
    def sheet_2_6(self):
        sources = [Funding.SourceType.FOREIGN, Funding.SourceType.STATE_PROGRAM,
                   Funding.SourceType.BUSINESS_CONTRACT]
        for row in iter_rows(self.wb["2,6"], 3, 5):
            t = self.teacher(row[1], "2,6")
            for col, source in zip((2, 3, 4), sources):
                amount = to_number(row[col])
                if t and amount:
                    Funding.objects.create(teacher=t, reporting_period=self.period, source_type=source,
                                           amount=amount, contract_info=IMPORT_NOTE)
                    self.stats["сумм средств (2,6)"] += 1

    # ----------------------------------------------------------------- 3,1
    def sheet_3_1(self):
        sources = [ReputationVote.Source.RATING_ORG, ReputationVote.Source.NATIONAL_QA,
                   ReputationVote.Source.FOREIGN_PARTNER]
        for row in iter_rows(self.wb["3,1"], 3, 5):
            t = self.teacher(row[1], "3,1")
            for col, source in zip((2, 3, 4), sources):
                votes = to_number(row[col])
                if t and votes:
                    ReputationVote.objects.create(teacher=t, reporting_period=self.period, source=source,
                                                  votes_count=int(votes), note=IMPORT_NOTE)
                    self.stats["голосов за репутацию (3,1)"] += 1

    # ----------------------------------------------------------------- 3,3
    def sheet_3_3(self):
        ranks = [Rank.TOP100, Rank.TOP300, Rank.TOP500, Rank.OTHER]
        for row in iter_rows(self.wb["3,3"], 3, 7):
            name = clean_text(row[1])
            student, _ = Student.objects.get_or_create(department=self.department, full_name=name)
            for col, rank in zip((2, 3, 4, 5), ranks):
                if to_number(row[col]):
                    StudentMobility.objects.create(
                        student=student, reporting_period=self.period,
                        program_type=StudentMobility.ProgramType.EXCHANGE, partner_rank=rank,
                    )
                    self.stats["мобильность студентов (3,3)"] += 1

    # ----------------------------------------------------------------- 3,4
    def sheet_3_4(self):
        for row in iter_rows(self.wb["3,4"], 3, 4):
            t = self.teacher(row[1], "3,4")
            for col, direction in ((2, TeacherMobility.Direction.OUTBOUND),
                                   (3, TeacherMobility.Direction.INBOUND)):
                for _ in range(int(to_number(row[col]) or 0) if t else 0):
                    TeacherMobility.objects.create(teacher=t, reporting_period=self.period,
                                                   direction=direction, order_info=IMPORT_NOTE)
                    self.stats["мобильность ППС (3,4)"] += 1


def clear_period_data(period):
    """Удаляет факты периода перед повторным импортом (сами карточки
    преподавателей и их файлы не трогает)."""
    for model in (TeacherPeriodSnapshot, Publication, TeacherHIndex, Patent, Funding,
                  ReputationVote, TeacherMobility, StudentMobility):
        model.objects.filter(reporting_period=period).delete()
