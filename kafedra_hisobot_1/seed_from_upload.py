"""
Разовый скрипт: загружает реальные данные преподавателей из присланного
файла (kafedra_hisobot_template.xlsx, лист "5" — штатная ведомость) в БД.
Используется здесь только чтобы проверить, что модели и exporter реально
работают на настоящих данных, а не на выдуманных. Запуск:

    python manage.py shell < seed_from_upload.py
"""
import datetime
import openpyxl

from apps.core.models import University, Faculty, Department
from apps.dictionaries.models import EmploymentType, Position, AcademicDegree, AcademicTitle, Country
from apps.staff.models import Teacher, TeacherPeriodSnapshot
from apps.workflow.models import ReportingPeriod

TEMPLATE = "apps/exporter/templates_store/kafedra_hisobot_template.xlsx"

uni, _ = University.objects.get_or_create(short_name="ФерГУ", defaults={"name": "Farg'ona Davlat Universiteti"})
fac, _ = Faculty.objects.get_or_create(university=uni, name="Amaliy matematika va informatika fakulteti", defaults={"short_name": "AMI"})
dep, _ = Department.objects.get_or_create(faculty=fac, name="Matematik analiz va differensial tenglamalar kafedrasi", defaults={"short_name": "MADT"})

period, _ = ReportingPeriod.objects.get_or_create(
    department=dep, period_label="2026-mart",
    defaults={"period_date": datetime.date(2026, 3, 31)},
)

et_asosiy, _ = EmploymentType.objects.get_or_create(code="asosiy", defaults={"name": "Асосий штат", "sort_order": 1})
et_ichki, _ = EmploymentType.objects.get_or_create(code="ichki_orindosh", defaults={"name": "Ички ўриндош", "sort_order": 2})
et_tashqi, _ = EmploymentType.objects.get_or_create(code="tashqi_orindosh", defaults={"name": "Ташқи ўриндош", "sort_order": 3})
et_soatbay, _ = EmploymentType.objects.get_or_create(code="soatbay", defaults={"name": "Соатбай", "sort_order": 4})

degree_map = {}
for code, name in [("fan_nomzodi", "Фан номзоди"), ("phd", "PhD"), ("fan_doktori", "Фан доктори"), ("dsc", "DSc")]:
    degree_map[code], _ = AcademicDegree.objects.get_or_create(code=code, defaults={"name": name})
# нормализация опечатки "Phd" -> "phd" делается ниже по .lower()

title_dotsent, _ = AcademicTitle.objects.get_or_create(name="Доцент")
title_professor, _ = AcademicTitle.objects.get_or_create(name="Профессор")

pos_cache = {}
def get_position(name):
    if not name:
        return None
    key = name.strip()
    if key not in pos_cache:
        pos_cache[key], _ = Position.objects.get_or_create(name=key)
    return pos_cache[key]

DEGREE_ALIASES = {
    "fan nomzodi": "fan_nomzodi",
    "phd": "phd",
    "dsc": "dsc",
    "fan doktori": "fan_doktori",
}
TITLE_ALIASES = {
    "dotsent": title_dotsent,
    "professor": title_professor,
}

wb = openpyxl.load_workbook(TEMPLATE, data_only=True)
ws = wb["5"]

SECTIONS = [
    (et_asosiy, range(7, 21)),
    (et_ichki, range(22, 23)),
    (et_tashqi, range(24, 25)),   # только строка с реальными данными
    (et_soatbay, range(28, 31)),
]

created = 0
for et, row_range in SECTIONS:
    for row in row_range:
        name = ws.cell(row=row, column=2).value
        if not name or not str(name).strip():
            continue
        name = str(name).strip()
        birth_year = ws.cell(row=row, column=3).value
        position_name = ws.cell(row=row, column=4).value
        stavka = ws.cell(row=row, column=5).value or 1
        degree_name = ws.cell(row=row, column=6).value
        title_name = ws.cell(row=row, column=7).value

        teacher, _ = Teacher.objects.get_or_create(
            department=dep, full_name=name,
            defaults={"birth_date": datetime.date(int(birth_year), 1, 1) if birth_year else None},
        )

        degree_obj = None
        if degree_name:
            code = DEGREE_ALIASES.get(str(degree_name).strip().lower())
            degree_obj = degree_map.get(code) if code else None

        title_obj = TITLE_ALIASES.get(str(title_name).strip().lower()) if title_name else None

        snap, _ = TeacherPeriodSnapshot.objects.update_or_create(
            teacher=teacher, reporting_period=period,
            defaults={
                "employment_type": et,
                "position": get_position(position_name),
                "stavka": stavka if isinstance(stavka, (int, float)) else 1,
                "academic_degree": degree_obj,
                "academic_title": title_obj,
            },
        )
        created += 1

print(f"Готово: кафедра={dep}, период={period}, загружено снимков={created}")

# ---------------------------------------------------------------------------
# Дозагрузка "постоянных" полей преподавателя (диплом, приказ, специальность)
# из листа "1,1" — сопоставление по полному имени (с обрезкой пробелов).
# ---------------------------------------------------------------------------
ws11 = wb["1,1"]


def norm(name):
    return " ".join(str(name).split()) if name else ""


teachers_by_name = {norm(t.full_name): t for t in Teacher.objects.filter(department=dep)}

DEG_COLS = [(4, 5), (6, 7), (8, 9)]

updated = 0
for row in list(range(4, 18)) + [19]:
    name = ws11.cell(row=row, column=2).value
    if not name or not str(name).strip():
        continue
    teacher = teachers_by_name.get(norm(name))
    if not teacher:
        continue

    teacher.foreign_university_name = ws11.cell(row=row, column=3).value or ""
    teacher.specialty_name = ws11.cell(row=row, column=10).value or ""
    order_raw = ws11.cell(row=row, column=11).value or ""
    if order_raw:
        import re
        parts = str(order_raw).split(",", 1)
        teacher.hire_order_number = parts[0].strip()
        if len(parts) > 1:
            m = re.search(r"(\d{1,2})\.(\d{1,2})\.(\d{4})", parts[1])
            if m:
                d, mo, y = map(int, m.groups())
                teacher.hire_order_date = datetime.date(y, mo, d)
    for c1, c2 in DEG_COLS:
        series = ws11.cell(row=row, column=c1).value
        number = ws11.cell(row=row, column=c2).value
        if series or number:
            teacher.diploma_series = str(series or "")
            teacher.diploma_number = str(number or "")
            break
    teacher.save()
    updated += 1

print(f"Дозаполнено постоянных полей преподавателя: {updated}")
